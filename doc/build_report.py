from pathlib import Path
from textwrap import wrap

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "doc" / "Z202508090307_杨文帅_实验四后端实现实验报告.docx"

ACCENT = "1B8278"
ACCENT_DARK = "123C39"
LIGHT = "EAF2F0"
PALE = "F5F8F7"
GRID = "D7E2DD"
TEXT = "202625"
MUTED = "687673"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRID, size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_font(run, name="Microsoft YaHei", size=10.5, bold=False, color=TEXT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, before=0, after=6, line=1.35, align=None):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def add_page_number(paragraph):
    run = paragraph.add_run("第 ")
    set_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    run2 = paragraph.add_run(" 页")
    set_font(run2, size=9, color=MUTED)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    style_paragraph(p, before=12 if level == 1 else 8, after=7 if level == 1 else 4, line=1.15)
    run = p.add_run(text)
    set_font(run, size=15 if level == 1 else 12, bold=True, color=TEXT)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    style_paragraph(p, after=6, line=1.38)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.55 + level * 0.35)
    p.paragraph_format.first_line_indent = Cm(-0.22)
    style_paragraph(p, after=3, line=1.25)
    r = p.add_run(text)
    set_font(r, size=10.2)
    return p


def add_table(doc, headers, rows, widths=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header = table.rows[0].cells
    for i, value in enumerate(headers):
        cell = header[i]
        set_cell_shading(cell, ACCENT_DARK)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        style_paragraph(p, after=0, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(str(value))
        set_font(r, size=font_size, bold=True, color="FFFFFF")
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(row):
            cell = cells[col_index]
            set_cell_shading(cell, "FFFFFF" if row_index % 2 == 0 else PALE)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            style_paragraph(p, after=0, line=1.12, align=WD_ALIGN_PARAGRAPH.CENTER if col_index == 0 else WD_ALIGN_PARAGRAPH.LEFT)
            r = p.add_run(str(value))
            set_font(r, size=font_size)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_code(doc, text, caption=None, max_chars=102):
    if caption:
        p = doc.add_paragraph()
        style_paragraph(p, before=5, after=3, line=1.0)
        r = p.add_run(caption)
        set_font(r, size=9.2, bold=True, color=ACCENT_DARK)
    box = doc.add_table(rows=1, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = box.cell(0, 0)
    set_cell_shading(cell, "F2F5F4")
    set_cell_border(cell, color=GRID, size="4")
    set_cell_margins(cell, top=120, start=160, bottom=120, end=160)
    lines = []
    for original in text.strip("\n").splitlines():
        if len(original) <= max_chars:
            lines.append(original)
        else:
            lines.extend(wrap(original, width=max_chars, replace_whitespace=False, drop_whitespace=False))
    p = cell.paragraphs[0]
    for index, line in enumerate(lines):
        if index:
            p.add_run().add_break()
        r = p.add_run(line)
        set_font(r, name="Consolas", size=7.7, color=ACCENT_DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return box


def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT)
    set_cell_border(cell, color=ACCENT, size="8")
    set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    p = cell.paragraphs[0]
    style_paragraph(p, after=3, line=1.2)
    r = p.add_run(title)
    set_font(r, size=10.5, bold=True, color=ACCENT_DARK)
    p2 = cell.add_paragraph()
    style_paragraph(p2, after=0, line=1.28)
    r = p2.add_run(text)
    set_font(r, size=9.8)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_figure(doc, path, caption, width=6.25):
    if not path.exists():
        add_callout(doc, "图表未找到", str(path))
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(str(path), width=Inches(width))
    style_paragraph(p, after=2, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    cap = doc.add_paragraph()
    style_paragraph(cap, after=6, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = cap.add_run(caption)
    set_font(r, size=9, color=MUTED)


def source(rel, start=None, end=None):
    lines = (ROOT / rel).read_text(encoding="utf-8").splitlines()
    if start is None:
        return "\n".join(lines)
    return "\n".join(lines[start - 1:end])


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.15)
    section.right_margin = Cm(2.15)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor.from_string(TEXT)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.text = "实验四 后端实现 | 全球天气数据分析与推荐系统"
    for run in header.runs:
        set_font(run, size=8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer)


def build():
    doc = Document()
    configure_document(doc)

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(55)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("实验四 后端实现实验报告")
    set_font(r, size=24, bold=True, color=TEXT)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(36)
    r = p.add_run("基于 Java Spring Boot 的全球天气数据分析与推荐系统")
    set_font(r, size=14, color=ACCENT_DARK)
    add_callout(doc, "项目说明", "本报告对应独立的 Spring Boot 天气项目，原 Python 天气项目保持不变。后端同时提供天气分析、登录、数据集管理、模型管理和模型预测接口，默认端口为 8091。")
    add_table(doc, ["项目字段", "内容"], [
        ["学生学号", "Z202508090307"],
        ["学生姓名", "杨文帅"],
        ["项目名称", "全球天气数据分析与推荐系统（Spring Boot 版本）"],
        ["后端技术", "Spring Boot 3.5.16、Java 17、MyBatis-Plus 3.5.17"],
        ["数据源", "GlobalWeatherRepository.csv，166253 条有效记录，41 个原始字段"],
        ["完成日期", "2026 年 10 月"],
    ], widths=[4.2, 11.0], font_size=10)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(65)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("项目目录：C:\\传媒项目\\_ex\\weather-system-springboot")
    set_font(r, size=9.2, color=MUTED)
    doc.add_page_break()

    add_heading(doc, "一、实验目的", 1)
    add_body(doc, "本实验围绕全球天气数据分析与推荐系统完成 Spring Boot 后端实现，目标是把数据处理、分析服务、管理接口和前端页面组织成可以独立运行、测试和重构的工程。")
    add_bullet(doc, "掌握 Spring Boot 项目搭建、Maven 依赖配置、application.yml 环境配置和 CORS 配置。")
    add_bullet(doc, "理解 Controller、Service、Mapper、Entity 和统一异常处理之间的职责边界。")
    add_bullet(doc, "完成用户登录、会话认证、数据集管理、模型管理和天气分析 RESTful 接口。")
    add_bullet(doc, "使用 MyBatis-Plus 连接 H2 文件数据库，并提供 MySQL 建库脚本和 profile 配置。")
    add_bullet(doc, "使用接口测试验证正常请求、未登录访问、参数校验和模型推理等场景。")
    add_heading(doc, "二、实验环境", 1)
    add_table(doc, ["类别", "环境或工具", "用途"], [
        ["操作系统", "Windows 10/11", "项目开发与运行"],
        ["Java", "JDK 17 或更高版本", "编译与运行 Spring Boot"],
        ["后端框架", "Spring Boot 3.5.16", "Web 服务、配置和静态资源"],
        ["持久层", "MyBatis-Plus 3.5.17", "实体映射、分页和 CRUD"],
        ["数据库", "H2 文件数据库；MySQL profile", "本地运行与部署演示"],
        ["构建工具", "Maven 3.9+", "依赖解析与打包"],
        ["Python 环境", "Python 3.14、pandas、NumPy、scikit-learn、LightGBM、joblib", "EDA、模型训练和推理桥接"],
        ["辅助工具", "VS Code、Postman/curl、Redis 5.0.14.1（可选）", "编辑、接口测试和扩展缓存"],
    ], widths=[3.0, 6.0, 6.2])
    add_callout(doc, "运行说明", "项目默认使用 H2，因此不配置 MySQL 也可以启动。MySQL 模式使用独立的 weather_springboot 数据库，不修改原电影项目数据库；原天气项目继续使用 8090，本项目使用 8091。")

    add_heading(doc, "三、项目概述与功能设计", 1)
    add_body(doc, "系统以 GlobalWeatherRepository.csv 为可替换数据源，读取国家、城市、坐标、温度、风速、气压、降水、湿度、能见度、紫外线和空气质量等字段，形成天气统计、城市检索、出行推荐、城市相似度、季节计划、EDA、降水预测和后台管理功能。")
    add_table(doc, ["模块", "主要功能", "对应接口或页面"], [
        ["用户认证", "登录、退出、当前用户、管理员会话校验", "/api/auth/login、/api/auth/me"],
        ["天气分析", "汇总、趋势、天气状况、地图、城市、推荐、对比、预测", "/api/summary、/api/trend、/api/recommend 等"],
        ["数据质量", "清洗统计、字段边界检查、数据集重载和 CSV 上传", "/api/cleaning、/api/datasets、/api/admin/reload"],
        ["模型实验", "八模型指标、图表、降水预测、风险城市排行", "/api/ml、/api/ml/predict、/api/ml/risk_rank"],
        ["后台管理", "用户、数据集、模型分页查询和增删改查", "/api/admin/users、/api/datasets、/api/models"],
    ], widths=[3.0, 7.0, 5.2])
    add_heading(doc, "四、项目搭建与目录结构", 1)
    add_body(doc, "新项目位于 weather-system-springboot，与原 weather-system 目录并列，复制了原有前端资源和分析图表，但后端改为 Java Spring Boot。项目采用按职责分包方式组织代码。")
    add_table(doc, ["目录", "职责"], [
        ["api", "WeatherController、AuthController、ModelPredictionController，负责 HTTP 请求与参数绑定。"],
        ["data", "CsvReader、WeatherDataService、WeatherAnalyticsService，负责数据读取、快照和分析。"],
        ["admin", "Entity、Mapper、数据库初始化和用户/数据集/模型 CRUD Controller。"],
        ["common", "统一异常处理、Python 模型训练任务和降水预测桥接。"],
        ["config", "BCrypt 密码编码、CORS 和 MyBatis-Plus 分页拦截器。"],
        ["resources/static", "登录页、天气分析前端、ECharts、世界地图和 EDA/模型图表。"],
        ["data/raw、algorithm、figures", "可替换 CSV、EDA/训练脚本、模型文件和原始分析图表。"],
    ], widths=[4.2, 11.0])
    add_code(doc, source("pom.xml", 25, 67), "代码 4-1 主要 Maven 依赖")
    add_code(doc, source("src/main/resources/application.yml", 1, 48), "代码 4-2 application.yml 的端口、数据库和数据文件配置")

    add_heading(doc, "五、数据库设计与持久层实现", 1)
    add_body(doc, "本项目默认使用 H2 文件数据库，启动时由 schema.sql 初始化表结构，DatabaseInitializer 自动创建 admin 管理员和初始数据集记录。为完成部署实验，项目同时提供 application-mysql.yml 和 sql/weather_springboot_mysql.sql。")
    add_table(doc, ["表名", "核心字段", "用途"], [
        ["app_user", "id、username、password_hash、display_name、role_name、created_at", "用户身份、密码摘要和管理员角色"],
        ["weather_dataset", "id、name、file_path、row_count、field_count、status、created_at、updated_at", "数据集登记、上传、激活和归档"],
        ["weather_model", "id、model_name、algorithm、target_name、status、metrics_json、artifact_path", "模型训练任务、指标和产物路径"],
    ], widths=[3.4, 8.0, 3.8])
    add_code(doc, source("src/main/resources/schema.sql", 1, 34), "代码 5-1 H2/MySQL 兼容建表 SQL")
    add_code(doc, "@Mapper\npublic interface UserMapper extends BaseMapper<UserEntity> {}\n\n@Mapper\npublic interface DatasetMapper extends BaseMapper<DatasetEntity> {}\n\n@Mapper\npublic interface ModelMapper extends BaseMapper<ModelEntity> {}", "代码 5-2 MyBatis-Plus Mapper")
    add_callout(doc, "分页配置", "项目补充 MybatisPlusConfig，在 MybatisPlusInterceptor 中注册 PaginationInnerInterceptor，保证后台 users、datasets 和 models 接口的 page/size 参数真正作用于 SQL 分页。")

    add_heading(doc, "六、代码实现", 1)
    add_heading(doc, "6.1 Controller 层", 2)
    add_body(doc, "WeatherController 只负责接收查询参数并调用分析服务，避免把统计逻辑堆积在控制器中。health 接口用于启动检查，summary、cities、recommend、eda 和 ml 接口为前端页面提供 JSON 数据。")
    add_code(doc, source("src/main/java/com/weatherlab/weather/api/WeatherController.java", 24, 50), "代码 6-1 天气分析 Controller")
    add_heading(doc, "6.2 登录与会话认证", 2)
    add_body(doc, "登录接口使用 BCryptPasswordEncoder 对密码进行匹配，验证成功后在 HttpSession 中写入 userId、username 和 role。管理员接口统一调用 AdminGuard.require，未登录或非管理员请求返回 401。密码摘要不会通过 JSON 返回。")
    add_code(doc, source("src/main/java/com/weatherlab/weather/api/AuthController.java", 28, 70), "代码 6-2 登录、退出和当前用户接口")
    add_heading(doc, "6.3 Service 层与 CSV 数据处理", 2)
    add_body(doc, "CsvReader 按表头建立列索引，校验必需字段，解析时间和数值字段，检查坐标范围，并用 country、city、updatedEpoch 组合指纹去除重复观测。WeatherDataService 以快照形式保存当前数据，上传新 CSV 通过校验后才激活。")
    add_code(doc, source("src/main/java/com/weatherlab/weather/data/WeatherDataService.java", 27, 58), "代码 6-3 数据加载和数据集激活")
    add_code(doc, source("src/main/java/com/weatherlab/weather/data/CsvReader.java", 29, 72), "代码 6-4 CSV 字段校验、去重和快照生成")
    add_heading(doc, "6.4 用户、数据集和模型 CRUD", 2)
    add_body(doc, "后台控制器使用 MyBatis-Plus QueryWrapper 和 Page 完成分页查询。创建和修改请求通过 Jakarta Validation 校验必填字段与长度，数据集上传仅允许 CSV，活动数据集不能直接删除，运行中的模型不能删除。")
    add_code(doc, source("src/main/java/com/weatherlab/weather/admin/DatasetAdminController.java", 52, 118), "代码 6-5 数据集登记、上传和重载接口")
    add_heading(doc, "6.5 参数校验与统一异常处理", 2)
    add_body(doc, "接口层使用 @Valid、@NotBlank、@Size、@Min、@Max、@DecimalMin 和 @DecimalMax 约束请求。ApiExceptionHandler 将校验失败、业务参数错误、资源不存在、权限错误和未知异常转换为统一 JSON 结构，便于前端显示和日志排查。")
    add_code(doc, source("src/main/java/com/weatherlab/weather/common/ApiExceptionHandler.java", 15, 48), "代码 6-6 统一异常响应")
    add_heading(doc, "6.6 模型训练与预测桥接", 2)
    add_body(doc, "模型训练由 ModelTrainingService 以后台任务启动 Python ml_train.py，状态写回 weather_model；在线预测由 RainPredictionService 调用 predict_cli.py，读取已训练的 LightGBM 模型并解析最后一行 JSON。模型运行环境缺失时接口返回 available=false，不影响基础天气分析页面。")
    add_code(doc, source("src/main/java/com/weatherlab/weather/api/ModelPredictionController.java", 21, 53), "代码 6-7 模型预测接口和输入边界")
    add_code(doc, source("src/main/java/com/weatherlab/weather/common/RainPredictionService.java", 48, 86), "代码 6-8 Java 到 Python 模型推理桥接")

    add_heading(doc, "七、数据分析与模型结果", 1)
    add_body(doc, "项目启动时读取 166253 条有效天气记录，包含 191 个国家或地区和 239 个城市位置。分析服务提供温度趋势、天气状况分布、空气质量、相关性、地理分布和生活指数。EDA 图表由 Python 脚本生成，Spring Boot 启动时从 static/figures 提供访问。")
    add_figure(doc, ROOT / "figures" / "01_temporal_coverage.png", "图 7-1 天气数据时间覆盖与每日城市覆盖")
    add_figure(doc, ROOT / "figures" / "06_feature_correlation.png", "图 7-2 天气与空气质量特征相关性")
    add_body(doc, "降水分类实验以 precip_mm > 0.1 作为标签，对比逻辑回归、决策树、随机森林、K 近邻、高斯朴素贝叶斯、支持向量机、梯度提升和 LightGBM 八个模型。现有训练报告中 LightGBM 的 F1 为 0.5724、AUC 为 0.9062，作为在线预测模型；图表和指标文件保存于 data/profile 与 figures。")
    add_table(doc, ["模型", "Accuracy", "Precision", "Recall", "F1", "AUC"], [
        ["逻辑回归", "0.8435", "0.5404", "0.1870", "0.2778", "0.8419"],
        ["决策树", "0.8599", "0.5911", "0.4211", "0.4918", "0.8503"],
        ["随机森林", "0.8795", "0.7299", "0.3994", "0.5163", "0.8972"],
        ["K 近邻", "0.8575", "0.6297", "0.2789", "0.3866", "0.8489"],
        ["高斯朴素贝叶斯", "0.6633", "0.3042", "0.8478", "0.4478", "0.8028"],
        ["支持向量机", "0.8651", "0.6962", "0.2876", "0.4070", "0.8370"],
        ["梯度提升", "0.8734", "0.7155", "0.3547", "0.4743", "0.8872"],
        ["LightGBM", "0.8845", "0.7085", "0.4801", "0.5724", "0.9062"],
    ], widths=[4.1, 2.1, 2.1, 2.1, 2.1, 2.1], font_size=8.8)
    add_figure(doc, ROOT / "figures" / "ml_model_comparison.png", "图 7-3 八种模型分类指标对比")
    add_figure(doc, ROOT / "figures" / "ml_confusion_matrix.png", "图 7-4 LightGBM 混淆矩阵")

    add_heading(doc, "八、接口测试与运行验证", 1)
    add_body(doc, "测试使用已打包的 weather-system-springboot-1.0.0.jar，在 127.0.0.1:8091 启动。服务日志显示 Tomcat、H2、MyBatis-Plus 和静态首页均初始化成功，CSV 加载完成后再执行接口测试。")
    add_table(doc, ["测试项", "请求", "结果"], [
        ["服务健康检查", "GET /api/health", "200；status=ok，records=166253"],
        ["登录成功", "POST /api/auth/login；admin/123456", "200；返回用户角色 ADMIN 和会话 Cookie"],
        ["数据汇总", "GET /api/summary", "200；返回字段数、国家数、时间范围、质量规则"],
        ["城市查询", "GET /api/cities?limit=3", "200；返回城市、国家、温度、湿度和空气质量"],
        ["推荐接口", "GET /api/recommend?limit=3", "200；按规则评分返回城市推荐"],
        ["EDA 接口", "GET /api/eda", "200；返回 conditions、trend、geospatial、correlation 等结构"],
        ["模型城市预测", "GET /api/ml/predict?k=China | Beijing", "200；LightGBM 可用，示例概率 2.2%"],
        ["模型手动预测", "POST /api/ml/predict", "200；示例输入返回 probability=5.1%"],
        ["管理员分页", "GET /api/admin/users?page=1&size=3", "200；返回 items、total、page、size"],
        ["未登录保护", "GET /api/datasets?page=1&size=1", "401；Please sign in first"],
        ["静态图表", "GET /figures/ml_model_comparison.png", "200；图表可被前端访问"],
    ], widths=[3.2, 6.3, 5.7], font_size=8.8)
    add_heading(doc, "九、环境问题与解决过程", 1)
    add_table(doc, ["问题", "原因", "解决方案"], [
        ["Maven 离线构建失败", "本地缓存记录的远程仓库 ID 不可用，且仓库需要写 tracking 文件", "改用联网 Maven 重新解析依赖，下载缺失的 Spring Boot 与 MyBatis-Plus 依赖。"],
        ["Map.of 编译错误", "Java Map.of 最多接收 10 对键值参数", "改用 LinkedHashMap 逐项 put，保持 JSON 字段顺序。"],
        ["分页拦截器缺失", "MyBatis-Plus 3.5.17 将分页依赖拆分", "增加 mybatis-plus-jsqlparser，并注册分页拦截器。"],
        ["Spring Boot 无法创建 Tomcat 临时目录", "系统 C:\\Temp 不可写", "在启动类中显式创建并指定 .runtime/docbase，同时将 JVM 临时目录指向 .runtime/tmp；start.cmd 也预先创建这两个目录。"],
        ["模型文件路径错误", "从 data/raw 路径回溯项目根目录少一层", "统一回溯到项目根目录后再定位 algorithm 和 logs。"],
        ["模型推理缺少 LightGBM", "joblib 文件由 LightGBM 生成，默认 Python 未安装包", "安装 lightgbm，并在 README 说明 PYTHON_COMMAND 与 requirements.txt。"],
        ["前端图表 500", "图表只在根目录 figures，未打包到 static", "复制图表到 src/main/resources/static/figures，并更新 run-eda.cmd 自动同步。"],
    ], widths=[4.0, 5.4, 5.8], font_size=8.8)
    add_callout(doc, "重构建议", "更换数据集时保留必要字段名或同步修改 CsvReader 的 required columns；先调用 /api/datasets/upload 验证并激活，再运行 EDA 和模型训练脚本。数据集采用快照缓存，失败上传不会替换当前有效数据。")

    add_heading(doc, "十、实验总结", 1)
    add_body(doc, "本次实验完成了一个独立的 Java Spring Boot 天气项目。通过分层结构将 HTTP 接口、天气分析、CSV 数据处理、MyBatis-Plus 持久化和 Python 模型推理连接起来，实现了从数据读取到前端展示的完整后端链路。")
    add_body(doc, "实验中重点解决了依赖解析、Java 17 兼容性、分页插件、临时目录权限、模型路径、Python 依赖和静态资源打包等环境问题。通过健康检查、登录、业务查询、管理分页、参数边界和模型预测测试，确认应用可以在 8091 端口稳定启动并提供主要功能。")
    add_body(doc, "后续重构可以继续引入 Redis 缓存城市推荐和热点查询，使用定时任务更新数据快照，增加 OpenAPI 文档与自动化测试，并把模型服务拆分为独立 Python 服务。当前项目已经把数据集、EDA、模型指标和前端资源放在清晰目录中，便于替换数据集和扩展业务模块。")
    add_callout(doc, "最终运行命令", "在项目目录执行 start.cmd，打开 http://127.0.0.1:8091/；默认账号 admin，密码 123456。打包命令为 mvn -DskipTests package，生成 target/weather-system-springboot-1.0.0.jar。")

    doc.core_properties.title = "实验四 后端实现实验报告"
    doc.core_properties.subject = "基于 Java Spring Boot 的全球天气数据分析与推荐系统"
    doc.core_properties.author = "杨文帅"
    doc.core_properties.keywords = "Spring Boot, MyBatis-Plus, Weather, Backend"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
